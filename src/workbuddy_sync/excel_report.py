from __future__ import annotations

from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .models import DownloadResult


REPORT_FILENAME = "专家专家团压缩包清单.xlsx"


HEADERS = ["压缩包名", "业务分类", "专家/专家团", "中文名", "plugin", "下载状态", "文件大小", "来源URL", "错误信息"]


def _status_zh(status: str) -> str:
    return {
        "success": "成功",
        "skipped": "已存在",
        "failed": "失败",
    }.get(status, status)


def _setup_sheet(ws) -> None:
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9E2F3")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=thin)


def _fit_columns(ws, min_width: int = 10, max_width: int = 60) -> None:
    for col_idx in range(1, ws.max_column + 1):
        letter = get_column_letter(col_idx)
        longest = 0
        for cell in ws[letter]:
            if cell.value is None:
                continue
            longest = max(longest, len(str(cell.value)))
        ws.column_dimensions[letter].width = min(max(longest + 2, min_width), max_width)


def _apply_template_widths(wb: Workbook, template_path: Path | None) -> None:
    if not template_path or not template_path.exists():
        return
    template = load_workbook(template_path, read_only=False, data_only=False)
    for sheet_name in ("专家", "专家团"):
        if sheet_name not in wb.sheetnames or sheet_name not in template.sheetnames:
            continue
        src = template[sheet_name]
        dst = wb[sheet_name]
        for col in ("A", "B", "C", "D"):
            width = src.column_dimensions[col].width
            if width:
                dst.column_dimensions[col].width = width


def _write_detail_sheet(wb: Workbook, title: str, results: list[DownloadResult]) -> None:
    ws = wb.create_sheet(title)
    ws.append(HEADERS)
    for result in results:
        entry = result.entry
        ws.append(
            [
                entry.output_filename,
                entry.category_name,
                entry.type_label,
                entry.profession_zh or entry.display_name_zh,
                entry.plugin,
                _status_zh(result.status),
                result.file_size,
                result.url,
                result.error,
            ]
        )
    _setup_sheet(ws)
    _fit_columns(ws)


def _write_failure_sheet(wb: Workbook, failed: list[DownloadResult]) -> None:
    ws = wb.create_sheet("失败重跑队列")
    ws.append(["压缩包名", "业务分类", "专家/专家团", "中文名", "plugin", "来源URL", "错误信息", "重试次数"])
    for result in failed:
        entry = result.entry
        ws.append(
            [
                entry.output_filename,
                entry.category_name,
                entry.type_label,
                entry.profession_zh or entry.display_name_zh,
                entry.plugin,
                result.url,
                result.error,
                result.retries,
            ]
        )
    _setup_sheet(ws)
    _fit_columns(ws)


def _write_dashboard(wb: Workbook, results: list[DownloadResult]) -> None:
    ws = wb.active
    ws.title = "统计看板"

    total = len(results)
    success = sum(1 for item in results if item.success)
    failed = total - success
    by_type = Counter(item.entry.type_label for item in results)

    ws["A1"] = "WorkBuddy 专家/专家团同步看板"
    ws["A1"].font = Font(size=16, bold=True, color="1F4E78")
    ws.merge_cells("A1:F1")

    summary_rows = [
        ("总条目数", total),
        ("下载成功/已存在", success),
        ("下载失败", failed),
        ("专家数量", by_type.get("专家", 0)),
        ("专家团数量", by_type.get("专家团", 0)),
    ]
    ws.append([])
    ws.append(["指标", "数值"])
    for label, value in summary_rows:
        ws.append([label, value])

    start_row = 10
    ws.cell(row=start_row, column=1, value="业务分类")
    ws.cell(row=start_row, column=2, value="专家")
    ws.cell(row=start_row, column=3, value="专家团")
    ws.cell(row=start_row, column=4, value="合计")

    category_counts: dict[str, Counter[str]] = defaultdict(Counter)
    for item in results:
        category_counts[item.entry.category_name][item.entry.type_label] += 1

    row = start_row + 1
    for category in sorted(category_counts):
        agent_count = category_counts[category].get("专家", 0)
        team_count = category_counts[category].get("专家团", 0)
        ws.cell(row=row, column=1, value=category)
        ws.cell(row=row, column=2, value=agent_count)
        ws.cell(row=row, column=3, value=team_count)
        ws.cell(row=row, column=4, value=agent_count + team_count)
        row += 1

    for header_row in (3, start_row):
        for cell in ws[header_row]:
            if cell.value is not None:
                cell.fill = PatternFill("solid", fgColor="1F4E78")
                cell.font = Font(color="FFFFFF", bold=True)
                cell.alignment = Alignment(horizontal="center")

    for row_cells in ws.iter_rows(min_row=4, max_row=8, min_col=1, max_col=2):
        row_cells[0].font = Font(bold=True)
        row_cells[1].alignment = Alignment(horizontal="right")

    _fit_columns(ws, max_width=36)


def write_report(
    results: list[DownloadResult],
    output_path: Path,
    template_path: Path | None = None,
) -> Path:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    _write_dashboard(wb, results)
    _write_detail_sheet(wb, "专家", [item for item in results if item.entry.type_label == "专家"])
    _write_detail_sheet(wb, "专家团", [item for item in results if item.entry.type_label == "专家团"])
    _write_failure_sheet(wb, [item for item in results if not item.success])
    _apply_template_widths(wb, template_path)
    wb.save(output_path)
    return output_path
